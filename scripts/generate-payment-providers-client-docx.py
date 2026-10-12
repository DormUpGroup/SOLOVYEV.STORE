"""Generate docs/PAYMENT_PROVIDERS_CLIENT.docx — Grow go-live checklist for the client."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "PAYMENT_PROVIDERS_CLIENT.docx"
OUT_FALLBACK = Path(__file__).resolve().parents[1] / "docs" / "PAYMENT_PROVIDERS_CLIENT_updated.docx"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "1C1C1E") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, header_fill)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
            for p in table.rows[ri].cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)


def save_doc(doc: Document) -> Path:
    try:
        doc.save(OUT)
        return OUT
    except PermissionError:
        doc.save(OUT_FALLBACK)
        return OUT_FALLBACK


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Подключение Grow — что сделать вам", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = [
        ("Для:", "SOLOVYEV STORE"),
        ("Сайт:", "solovyev.store"),
        ("Дата:", "28 сентября 2026"),
        ("Статус:", "сайт уже готов к Grow. Осталось открыть счёт и прислать доступы."),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(label + " ").bold = True
        p.add_run(value)

    doc.add_paragraph(
        "Оплата картой на сайте сделана. Покупатель жмёт оплатить, уходит на страницу Grow, "
        "платит картой (или Bit, если это включено в вашем тарифе) и возвращается на solovyev.store. "
        "Карту мы не храним."
    )

    doc.add_heading("1. Что уже готово на сайте", level=1)
    add_bullets(
        doc,
        [
            "Корзина → /checkout → страница Grow → возврат на /checkout/success",
            "Заказ SS-… создаётся до оплаты",
            "После успешной оплаты статус сам становится Paid",
            "В админке видны сумма, доставка, last4 карты и id транзакции Grow",
            "Чек (קבלה) выставляет Grow. Магазин — עוסק פטור, без НДС, не חשבונית מס",
        ],
    )
    doc.add_paragraph(
        "Пока нет ключей Grow, на checkout будет сообщение, что оплата ещё не настроена. "
        "Витрина при этом работает."
    )

    doc.add_heading("2. Что нужно от вас (коротко)", level=1)
    add_numbered(
        doc,
        [
            "Открыть счёт в Grow (раньше компания называлась Meshulam).",
            "Прислать нам два кода: UserId и PageCode.",
            "Согласовать с Grow адрес уведомлений (callback) — см. ниже.",
            "Сделать одну тестовую оплату на ₪1–5 и прислать номер заказа SS-….",
        ],
    )
    doc.add_paragraph(
        "После этого мы вставляем коды на сервер и включаем оплату на боевом сайте."
    )

    doc.add_heading("3. Как открыть Grow", level=1)
    add_numbered(
        doc,
        [
            "Зайти на grow.me / кабинет Meshulam и подать заявку как עוסק פטור.",
            "Обычно счёт открывают за 1–3 дня, если документы в порядке.",
            "В кабинете создать страницу оплаты (payment page) для интернет-магазина.",
            "В настройках страницы указать сайт, успешный возврат, callback, валюту ILS и без НДС.",
        ],
    )
    add_table(
        doc,
        ["Поле", "Значение"],
        [
            ["Сайт", "https://solovyev.store"],
            ["Успешный возврат", "https://solovyev.store/checkout/success"],
            ["Уведомление серверу (notify / IPN / webhook)", "https://solovyev.store/api/payments/grow/callback"],
            ["Валюта", "ILS (₪)"],
            ["НДС", "нет (עוסק פטור)"],
        ],
    )
    add_note(
        doc,
        "Если менеджер Grow просит «success URL с номером заказа» — можно оставить общий "
        "/checkout/success. Номер заказа сайт подставляет сам.",
    )

    doc.add_heading("4. Что прислать разработчику", level=1)
    doc.add_paragraph(
        "Пришлите одним сообщением (WhatsApp / почта), не в общий чат с клиентами:"
    )
    add_bullets(
        doc,
        [
            "UserId",
            "PageCode",
            "это sandbox (тест) или боевой кабинет",
            "скрин или текст, что callback указан как https://solovyev.store/api/payments/grow/callback",
        ],
    )
    doc.add_paragraph("Этого достаточно. Логины кабинета Grow нам не нужны.")

    doc.add_heading("5. Как это выглядит для покупателя", level=1)
    add_numbered(
        doc,
        [
            "Входит в аккаунт на сайте.",
            "Указывает израильский мобильный 05xxxxxxxx — Grow требует его для 3D Secure.",
            "Выбирает доставку и жмёт оплатить.",
            "Открывается страница Grow. Там карта (и Bit, если включён у вас).",
            "После оплаты возвращается на сайт. Заказ в админке — Paid.",
        ],
    )
    doc.add_paragraph(
        "Деньги на ваш счёт Grow обычно отдаёт с 1-го числа следующего месяца. "
        "Часть суммы иногда можно вывести раньше — это условие в договоре Grow, не на сайте."
    )
    add_note(
        doc,
        "Ориентир цены Grow: ₪59–69/мес + 0.75% по израильской карте, иностранная карта дороже (~3.5%). "
        "Точная цифра — только в вашей оферте.",
    )

    doc.add_heading("6. Что делаем мы после ваших кодов", level=1)
    add_numbered(
        doc,
        [
            "Пишем ключи в Vercel (боевой сайт) и в локальный .env.local для проверки.",
            "Для теста: GROW_API_BASE=https://sandbox.meshulam.co.il/api/light/server/1.0",
            "Для боя: GROW_API_BASE=https://secure.meshulam.co.il/api/light/server/1.0",
            "Проверяем, что база знает поля Grow (миграция 014_grow_checkout.sql).",
            "Делаем тестовую оплату. Grow должен получить от сайта ответ 200 и подтверждение транзакции (approveTransaction). Без этого боевые ключи часто не выдают.",
        ],
    )

    doc.add_heading("7. Чеклист перед запуском", level=1)
    add_bullets(
        doc,
        [
            "Счёт Grow открыт",
            "Есть UserId и PageCode",
            "В кабинете Grow прописан callback https://solovyev.store/api/payments/grow/callback",
            "Коды переданы разработчику",
            "Тестовая оплата прошла, заказ в админке Paid",
            "Grow подтвердил интеграцию и дал боевые ключи (если тест был в sandbox)",
            "На Vercel стоят боевые GROW_USER_ID, GROW_PAGE_CODE, GROW_API_BASE",
            "После смены ключей сайт передеплоен",
        ],
    )

    doc.add_heading("8. Важно понимать", level=1)
    add_bullets(
        doc,
        [
            "Карту покупателя хранит Grow, не сайт.",
            "Пока ключей нет, оплатить на сайте нельзя — это нормально.",
            "Старые заказы через WhatsApp в админке остаются. Новые идут через карту.",
            "WhatsApp по-прежнему нужен для Sell / Trade и вопросов, не для оплаты.",
            "Договор, комиссии и срок выплаты денег — между вами и Grow. Сайт к этому не подключается.",
        ],
    )

    saved = save_doc(doc)
    print(f"Wrote {saved}")


if __name__ == "__main__":
    main()
