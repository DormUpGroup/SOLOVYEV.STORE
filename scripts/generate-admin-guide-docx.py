"""Generate docs/ADMIN_GUIDE.docx — admin panel memo for SOLOVYEV STORE."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "ADMIN_GUIDE.docx"
OUT_FALLBACK = Path(__file__).resolve().parents[1] / "docs" / "ADMIN_GUIDE_updated.docx"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    header_fill: str = "1C1C1E",
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, header_fill)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(10)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
            for p in table.rows[ri].cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)


def save_doc(doc: Document) -> Path:
    try:
        doc.save(OUT)
        return OUT
    except PermissionError:
        doc.save(OUT_FALLBACK)
        return OUT_FALLBACK


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Памятка для администраторов SOLOVYEV STORE", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run(
        "Краткая инструкция по работе с админ-панелью. "
        "Клиент оформляет заказ на сайте и платит картой через Grow. "
        "WhatsApp остаётся для Sell/Trade и вопросов."
    )

    # 1
    doc.add_heading("1. Вход", level=1)
    add_numbered(
        doc,
        [
            "Откройте секретный URL админки (его даёт разработчик; публичный /admin для посторонних не открывается).",
            "На экране входа нажмите Sign in и введите логин и пароль.",
            "После входа слева — меню по группам, сверху — Refresh, Publish to site и внизу sidebar — View store / Log out.",
        ],
    )
    doc.add_paragraph(
        "Не передавайте ссылку, логин и пароль клиентам. "
        "Сессия админа отдельна от личного кабинета покупателя (/account)."
    )
    doc.add_paragraph("На телефоне меню открывается кнопкой ☰ слева вверху.")

    # 2
    doc.add_heading("2. Карта разделов", level=1)
    doc.add_paragraph("Меню слева разбито на 4 группы:")

    doc.add_heading("Store — товары", level=2)
    add_table(
        doc,
        ["Раздел", "Зачем"],
        [
            ["Catalog", "Основной каталог"],
            ["Made to order", "Товары под заказ (страница /made-to-order)"],
            ["Brand new", "Новинки (страница /brand-new)"],
        ],
    )

    doc.add_heading("Commerce — заказы и клиенты", level=2)
    add_table(
        doc,
        ["Раздел", "Зачем"],
        [
            ["Orders", "Заказы с сайта, статусы, оплата"],
            ["Users", "Клиенты (аккаунты на сайте)"],
            ["Reporting", "Отчёты, графики, PDF по продажам"],
        ],
    )

    doc.add_heading("Content — тексты и настройки", level=2)
    add_table(
        doc,
        ["Раздел", "Зачем"],
        [
            ["FAQ", "Вопросы–ответы на сайте"],
            ["Settings", "WhatsApp, Instagram, hero, контакты"],
        ],
    )

    doc.add_heading("Insights — аналитика", level=2)
    add_table(
        doc,
        ["Раздел", "Зачем"],
        [
            ["Overview", "Быстрая сводка: товары, GA, последние события"],
            ["Analytics", "GA4 + воронка магазина + график тренда"],
        ],
    )

    p = doc.add_paragraph()
    p.add_run("Важно: ").bold = True
    p.add_run(
        "после правок каталога, FAQ или Settings нажмите Publish to site, "
        "чтобы изменения быстрее появились на витрине."
    )

    # 3
    doc.add_heading("3. Каталог (типичный день)", level=1)
    add_numbered(
        doc,
        [
            "Откройте Catalog (или Made to order / Brand new).",
            "Нажмите Add product (или Add made-to-order item / Add brand-new item).",
            "Заполните название, цену, размеры, загрузите фото.",
            "Выставьте статус: available / new_drop — можно заказывать; reserved — временно занят; sold — продан; draft — черновик, покупатели не видят.",
            "Нажмите Save в форме товара.",
            "Нажмите Publish to site сверху.",
        ],
    )

    doc.add_heading("Фото", level=2)
    add_bullets(
        doc,
        [
            "Загружаются прямо из карточки товара.",
            "Можно менять порядок (drag), кроп, поворот.",
            "Первое фото = главное на витрине.",
            "Справа есть Live preview — как будет выглядеть на сайте.",
        ],
    )

    doc.add_heading("Поиск и фильтры", level=2)
    add_bullets(
        doc,
        [
            "Строка Search products… и фильтр по бренду.",
            "Товары сгруппированы по бренду и категории.",
        ],
    )

    # 4
    doc.add_heading("4. Как клиент делает заказ", level=1)
    add_numbered(
        doc,
        [
            "Клиент входит в аккаунт на сайте.",
            "Жмёт checkout / Buy now.",
            "На /checkout выбирает доставку и платит картой.",
            "Система создаёт заказ SS-…, Grow списывает карту, статус становится Paid.",
        ],
    )
    doc.add_paragraph(
        "Если клиент не залогинен — его отправят на вход. "
        "Если оплата не настроена (нет ключей Grow) — на checkout будет ошибка, "
        "витрина работает как раньше."
    )

    doc.add_heading("Включить оплату (Grow)", level=2)
    add_numbered(
        doc,
        [
            "Применить миграции 013_payplus_checkout.sql (статусы pending_payment) и 014_grow_checkout.sql. Только SQL: node scripts/apply-migration.mjs 014_grow_checkout.sql.",
            "Открыть счёт Grow (UserId + PageCode). Для разработки — sandbox.",
            "Callback URL в кабинете Grow: https://solovyev.store/api/payments/grow/callback",
            "В .env.local / Vercel: GROW_USER_ID, GROW_PAGE_CODE, GROW_API_BASE. После смены ключей — Redeploy.",
            "Магазин — עוסק פטור: без НДС, клиент получает קבלה, не חשבונית מס.",
            "Grow проверяет интеграцию (callback 200 + approveTransaction) и только потом даёт боевые ключи.",
        ],
    )
    add_note(doc, "Чеклист для клиента: docs/PAYMENT_PROVIDERS_CLIENT.md")

    # 5
    doc.add_heading("5. Orders — ведение заказа", level=1)

    doc.add_heading("Список", level=2)
    add_bullets(
        doc,
        [
            "Поиск по номеру SS-… или email.",
            "Фильтр по статусу.",
            "Переключатель 7 / 30 / 90 days — сводка и график.",
            "Клик по строке → карточка заказа внизу.",
        ],
    )

    doc.add_heading("Карточка заказа", level=2)
    add_bullets(
        doc,
        [
            "Клиент, состав, сумма, доставка, last4 карты, id транзакции Grow.",
            "Смена статуса (сохраняется сразу).",
            "Copy order ref — скопировать SS-….",
            "Copy reply / Go to chat — если нужно написать клиенту в WhatsApp.",
        ],
    )

    doc.add_heading("Статусы (рекомендуемый порядок)", level=2)
    doc.add_paragraph(
        "Pending payment → Paid → Shipped → Completed (или Cancelled)"
    )
    add_table(
        doc,
        ["Статус", "Когда ставить"],
        [
            ["Pending payment", "Заказ создан, карта ещё не списана"],
            ["Pending WhatsApp", "Старые заказы до оплаты на сайте"],
            ["In chat", "Переписка (редко, для старых заказов)"],
            ["Paid", "Оплата прошла (ставится автоматически после Grow)"],
            ["Shipped", "Отправлено или передано клиенту"],
            ["Completed", "Сделка закрыта"],
            ["Cancelled", "Отмена, дубль, нет товара"],
        ],
    )

    doc.add_heading("Чеклист менеджера", level=2)
    add_numbered(
        doc,
        [
            "Дождаться Paid (ставится сам после Grow) или открыть заказ, если клиент написал в WhatsApp.",
            "В Catalog поставить reserved / sold при необходимости.",
            "Отгрузить → Shipped, закрыть → Completed.",
            "Если сорвалось → Cancelled, вернуть товар в available.",
        ],
    )
    add_note(
        doc,
        "Смена статуса заказа не меняет автоматически статус товара в каталоге.",
    )

    doc.add_heading("График в Orders", level=2)
    add_bullets(
        doc,
        [
            "Линия Orders — количество заказов по дням.",
            "Линия Paid revenue — выручка только по заказам со статусом Paid.",
        ],
    )

    # 6
    doc.add_heading("6. Reporting — отчёты", level=1)
    doc.add_paragraph("Раздел для обзора продаж без работы с отдельными заказами.")
    doc.add_heading("Что внутри", level=2)
    add_bullets(
        doc,
        [
            "Переключатель 7 / 30 / 90 days.",
            "Карточки: все заказы, выручка (paid), средний чек, число проданных позиций.",
            "Orders by status — круговая диаграмма: сколько заказов на каждом этапе.",
            "Top products by revenue — топ товаров по выручке (paid).",
            "Таблица топ-товаров с количеством и суммой.",
            "Download PDF — скачать PDF-отчёт за выбранный период.",
        ],
    )
    p = doc.add_paragraph()
    p.add_run("Revenue (paid)").bold = True
    p.add_run(" — только заказы со статусом Paid, не сумма всех корзин.")

    # 7
    doc.add_heading("7. Users — клиенты", level=1)
    add_bullets(
        doc,
        [
            "Поиск по email или имени.",
            "Фильтр Has orders — только с заказами.",
            "В карточке: контакты, marketing opt-in, заказы, избранное, корзина.",
        ],
    )
    doc.add_paragraph(
        "Телефон клиент может указать в /account — он нужен для 3D Secure при оплате картой."
    )
    p = doc.add_paragraph()
    p.add_run("Marketing opt-in").bold = True
    p.add_run(" — согласие на рекламные письма. Без него не используйте email для рассылок.")

    # 8
    doc.add_heading("8. Settings (осторожно)", level=1)
    add_bullets(
        doc,
        [
            "WhatsApp phone и manager name — ссылки и текст в сообщениях заказа.",
            "Instagram, hero-фото, название магазина.",
            "После правок — Publish to site.",
        ],
    )
    doc.add_paragraph("Неверный номер WhatsApp = клиенты уходят «в никуда».")

    # 9
    doc.add_heading("9. Analytics — аналитика", level=1)

    doc.add_heading("Google Analytics (если настроен)", level=2)
    add_bullets(
        doc,
        [
            "Realtime users, sessions, page views.",
            "Источники трафика и топ страниц.",
            "Смотрите Overview (кратко) или Analytics (подробно).",
        ],
    )

    doc.add_heading("Store funnel — воронка магазина", level=2)
    doc.add_paragraph("Три шага на сайте:")
    add_numbered(
        doc,
        [
            "Product views — просмотры товара",
            "Add to cart — добавления в корзину",
            "Checkout — переход к оплате на сайте",
        ],
    )
    doc.add_paragraph(
        "График-воронка показывает, на каком шаге теряются клиенты. "
        "В tooltip — % конверсии между шагами."
    )

    doc.add_heading("Trend", level=2)
    doc.add_paragraph(
        "График просмотров / корзины / checkout по дням за выбранный период (7 / 30 / 90 days)."
    )
    add_note(
        doc,
        "Для операционной работы по заказам в первую очередь смотрите Orders и Reporting, а не Analytics.",
    )

    # 10
    doc.add_heading("10. FAQ", level=1)
    add_bullets(
        doc,
        [
            "Редактирование вопросов и ответов в таблице.",
            "+ Add FAQ → заполнить → Save FAQ.",
            "После сохранения FAQ публикуется автоматически; для других разделов нужен Publish to site.",
        ],
    )

    # 11
    doc.add_heading("11. Частые ситуации", level=1)
    add_table(
        doc,
        ["Ситуация", "Что делать"],
        [
            [
                "Клиент пишет в WA, заказа нет",
                "Попросить оформить через сайт после входа в аккаунт",
            ],
            [
                "Чат не находится по SS-…",
                "Клиент ещё не отправил сообщение, или неверный номер в Settings",
            ],
            ["Дубль заказа", "Один оставить, второй → Cancelled"],
            [
                "Продали офлайн",
                "Catalog → sold, открытые заказы закрыть/отменить",
            ],
            [
                "На сайте «старое» после правок",
                "Publish to site → подождать минуту → Ctrl+F5",
            ],
            [
                "Много pending, мало paid",
                "Reporting → Orders by status, работать с висящими заказами в Orders",
            ],
            [
                "Checkout: Card payments are not configured",
                "Нет GROW_USER_ID / GROW_PAGE_CODE на сервере",
            ],
            [
                "Заказ завис в Pending payment",
                "Клиент не закончил оплату на странице Grow; можно отменить и вернуть товар в available",
            ],
            [
                "Оплата прошла, статус не Paid",
                "Проверить callback URL в кабинете Grow и id транзакции в карточке заказа",
            ],
        ],
    )

    # 12
    doc.add_heading("12. Чего в админке пока нет", level=1)
    add_bullets(
        doc,
        [
            "WhatsApp-чат внутри админки (переписка только в приложении WhatsApp)",
            "Автосмена статуса товара при смене статуса заказа",
            "Массовые рассылки и экспорт клиентов в CSV",
            "Заметки по клиенту / трекинг-код в UI (поля в базе есть, экран — позже)",
        ],
    )

    # 13
    doc.add_heading("13. Google Analytics — настройка (для разработчика)", level=1)
    doc.add_paragraph("Если GA не показывается, на сервере нужны:")
    add_bullets(
        doc,
        [
            "GA4_PROPERTY_ID — цифровой Property ID, не G-…",
            "GA4_CLIENT_EMAIL — email service account",
            "GA4_PRIVATE_KEY — private_key из JSON-ключа",
        ],
    )
    doc.add_paragraph(
        "Service account добавить в GA4 property с ролью Viewer, "
        "включить Google Analytics Data API."
    )

    # 14
    doc.add_heading("14. Безопасность", level=1)
    add_bullets(
        doc,
        [
            "Не передавать админ-URL и пароль.",
            "Log out на чужих компьютерах.",
            "Статусы заказов менять в Orders, а не вручную в базе.",
        ],
    )

    # 15
    doc.add_heading("15. Если что-то сломалось", level=1)
    doc.add_paragraph(
        "Напишите разработчику: что делали, время, номер заказа / email клиента, скрин ошибки."
    )
    doc.add_paragraph(
        "Типичные проблемы: не логинится, не создаются заказы SS-…, "
        "Publish не обновляет сайт, не загружаются фото, оплата не переводит заказ в Paid."
    )

    path = save_doc(doc)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
